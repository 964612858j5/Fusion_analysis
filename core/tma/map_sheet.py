"""TMA map sheet (F1, docs/tma_foundation_application.md §1, §4 F1): read a
TMA designer workbook (`TMA138.x.xlsx`) and check it strictly.

The workbook has four sheets:

* ``Information`` -- key / value rows: block name, layout name, layout size
  ``"<rows>x<cols>"``, core diameter, number of donors and of cores.
* ``<block id>`` (e.g. ``RB_1``) -- the layout grid: a header row ``X1..Xn``,
  a first column ``Y1..Ym``, and in each OCCUPIED cell its position id
  ``(X - 1) * rows + Y`` (every column numbered top to bottom).
* ``<block id>_Donor`` -- the same grid with the donor block id per cell.
* ``Lookup`` -- one row per LAYOUT position 1..rows*cols, also the empty
  ones: Unique TMA sample ID, TMA Block ID, Donor Block ID (empty for an
  empty position), Donor Core Height, Annotation Name and, in some
  workbooks, Spot Slide ``"<col>_(<row>;<col>)_<donor>"``.

Three kinds of check, kept apart so that a valid workbook is never refused
(codex review 2026-10-16): the LAYOUT (size, numbering formula), the
OCCUPIED positions (the three sheets agree on ids and donors, the count
matches Information) and OPTIONAL fields (Spot Slide checked only where
present). Every problem is collected with its sheet / cell; any problem
raises `MapSheetError` listing them all.

The donor block id is a tissue block, NOT a patient id; nothing here reads
it as one. How the grid lies on the slide (rotation, flip) is the
dearray's business (F2), not this module's.

The parser is general (any rows x cols the workbook declares); 19 x 8 is
the layout of the TMA138 blocks at hand, not a rule of this module.
"""

from __future__ import annotations

import hashlib
import os
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

SPOT = re.compile(r"^\s*(\d+)_\((\d+);(\d+)\)_(.+?)\s*$")
LAYOUT_SIZE = re.compile(r"^\s*(\d+)\s*[xX×]\s*(\d+)\s*$")
GRID_X = re.compile(r"^\s*X(\d+)\s*$")
GRID_Y = re.compile(r"^\s*Y(\d+)\s*$")


class MapSheetError(ValueError):
    """The workbook is not a valid map sheet; `problems` lists why."""

    def __init__(self, path: str, problems: List[str]):
        self.path = path
        self.problems = list(problems)
        super().__init__(f"{os.path.basename(path)}: " + "; ".join(self.problems))


@dataclass(frozen=True)
class Position:
    """One layout position (occupied or empty)."""
    position_id: int
    x: int                      # column, 1-based (X1..Xn)
    y: int                      # row, 1-based (Y1..Ym)
    donor_block_id: str = ""    # "" = empty position
    sample_id: str = ""         # "Unique TMA sample ID"
    core_height: str = ""
    annotation: str = ""
    spot_slide: str = ""

    @property
    def occupied(self) -> bool:
        return bool(self.donor_block_id)

    @property
    def label(self) -> str:
        return f"X{self.x}-Y{self.y:02d}"


@dataclass
class MapSheet:
    path: str
    sha256: str
    size: int
    block_id: str
    block_name: str
    rows: int
    cols: int
    layout_name: str
    diameter_mm: Optional[float]
    declared_cores: Optional[int]
    declared_donors: Optional[int]
    positions: Dict[int, Position] = field(default_factory=dict)
    has_spot_slide: bool = False

    def occupied(self) -> List[Position]:
        return [p for _, p in sorted(self.positions.items()) if p.occupied]

    def to_dict(self) -> Dict:
        return {
            "path": self.path, "sha256": self.sha256, "size": self.size,
            "block_id": self.block_id, "block_name": self.block_name,
            "rows": self.rows, "cols": self.cols, "layout_name": self.layout_name,
            "diameter_mm": self.diameter_mm, "declared_cores": self.declared_cores,
            "declared_donors": self.declared_donors, "has_spot_slide": self.has_spot_slide,
            "positions": [{"position_id": p.position_id, "x": p.x, "y": p.y,
                           "donor_block_id": p.donor_block_id, "sample_id": p.sample_id,
                           "core_height": p.core_height, "annotation": p.annotation,
                           "spot_slide": p.spot_slide}
                          for _, p in sorted(self.positions.items())],
        }


def position_id(x: int, y: int, rows: int) -> int:
    """The workbook's numbering: every column top to bottom."""
    return (x - 1) * rows + y


def file_sha256(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _int(v) -> Optional[int]:
    """A whole number, else None: 1.9 or "1.5" is NOT read as 1 (codex
    review), nor is inf / nan; callers report the original value."""
    if isinstance(v, bool):
        return None
    if isinstance(v, int):
        return v
    if isinstance(v, float):
        return int(v) if v == v and abs(v) != float("inf") and v.is_integer() else None
    s = _text(v)
    return int(s) if re.fullmatch(r"[+-]?\d+", s) else None


def _rows(ws) -> List[Tuple]:
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def _information(ws, problems) -> Dict[str, str]:
    info = {}
    for r in _rows(ws):
        cells = [_text(c) for c in r]
        if len(cells) >= 2 and cells[0]:
            info[cells[0].lower()] = cells[1]
    return info


def _grid(ws, name, problems) -> Tuple[Dict[Tuple[int, int], str], int, int]:
    """``{(x, y): text}`` of the occupied cells and the (rows, cols) its
    header and first column declare."""
    rows = _rows(ws)
    if not rows:
        problems.append(f"{name}: empty sheet")
        return {}, 0, 0
    header = [_text(c) for c in rows[0]]
    xs = {}
    for ci, h in enumerate(header[1:], start=1):
        m = GRID_X.match(h)
        if m:
            xs[ci] = int(m.group(1))
        elif h:
            problems.append(f"{name}!R1C{ci + 1}: header {h!r} is not X<n>")
    cells, ys = {}, []
    for ri, r in enumerate(rows[1:], start=2):
        first = _text(r[0]) if r else ""
        if not first and not any(_text(c) for c in r):
            continue
        m = GRID_Y.match(first)
        if not m:
            problems.append(f"{name}!R{ri}C1: row label {first!r} is not Y<n>")
            continue
        y = int(m.group(1))
        ys.append(y)
        for ci, x in xs.items():
            v = _text(r[ci]) if ci < len(r) else ""
            if v:
                cells[(x, y)] = v
    n_rows, n_cols = (max(ys) if ys else 0), (max(xs.values()) if xs else 0)
    if sorted(xs.values()) != list(range(1, n_cols + 1)):
        problems.append(f"{name}: columns are not X1..X{n_cols}")
    if sorted(ys) != list(range(1, n_rows + 1)):
        problems.append(f"{name}: rows are not Y1..Y{n_rows}")
    return cells, n_rows, n_cols


def _lookup(ws, problems) -> Tuple[List[Dict[str, str]], Dict[str, int]]:
    rows = _rows(ws)
    if not rows:
        problems.append("Lookup: empty sheet")
        return [], {}
    header = {(_text(h) or f"#{i}").lower(): i for i, h in enumerate(rows[0])}
    need = ["unique tma sample id", "tma block id", "donor block id"]
    for k in need:
        if k not in header:
            problems.append(f"Lookup: column {k!r} missing")
    out = []
    for ri, r in enumerate(rows[1:], start=2):
        rec = {k: (_text(r[i]) if i < len(r) else "") for k, i in header.items()}
        if not any(rec.values()):
            continue
        rec["_row"] = ri
        out.append(rec)
    return out, header


def parse(path: str) -> MapSheet:
    """Read and check the workbook at `path`; raises `MapSheetError`."""
    import openpyxl
    problems: List[str] = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:                       # not a workbook at all
        raise MapSheetError(path, [f"cannot open: {exc}"])
    try:
        return _parse(path, wb, problems)
    finally:
        wb.close()


def _parse(path, wb, problems) -> MapSheet:
    names = wb.sheetnames
    for need in ("Lookup", "Information"):
        if need not in names:
            problems.append(f"sheet {need!r} missing")
    if problems:
        raise MapSheetError(path, problems)
    lookup, header = _lookup(wb["Lookup"], problems)
    info = _information(wb["Information"], problems)
    block_ids = sorted({r.get("tma block id", "") for r in lookup if r.get("tma block id")})
    if len(block_ids) != 1:
        problems.append(f"Lookup: expected one TMA Block ID, found {block_ids}")
        raise MapSheetError(path, problems)
    block = block_ids[0]
    for need in (block, f"{block}_Donor"):
        if need not in names:
            problems.append(f"sheet {need!r} missing")
    if problems:
        raise MapSheetError(path, problems)
    ids, g_rows, g_cols = _grid(wb[block], block, problems)
    donors, d_rows, d_cols = _grid(wb[f"{block}_Donor"], f"{block}_Donor", problems)

    # ---- layout ----
    size = LAYOUT_SIZE.match(info.get("layout size (row x column)", ""))
    rows, cols = (int(size.group(1)), int(size.group(2))) if size else (g_rows, g_cols)
    if not size:
        problems.append("Information: 'Layout size (row x column)' missing or unreadable")
    if (g_rows, g_cols) != (rows, cols):
        problems.append(f"{block}: grid is {g_rows}x{g_cols}, Information says {rows}x{cols}")
    if (d_rows, d_cols) != (rows, cols):
        problems.append(f"{block}_Donor: grid is {d_rows}x{d_cols}, Information says {rows}x{cols}")
    for (x, y), v in sorted(ids.items()):
        n = _int(v)
        if n is None:
            problems.append(f"{block}!X{x}Y{y}: {v!r} is not a number")
        elif n != position_id(x, y, rows):
            problems.append(f"{block}!X{x}Y{y}: id {n}, the layout numbering gives {position_id(x, y, rows)}")
    n_layout = rows * cols

    # ---- occupied positions: the three sheets agree ----
    by_id: Dict[int, Dict[str, str]] = {}
    for rec in lookup:
        n = _int(rec.get("unique tma sample id"))
        where = f"Lookup!R{rec['_row']}"
        if n is None:
            problems.append(f"{where}: Unique TMA sample ID {rec.get('unique tma sample id')!r} is not a number")
            continue
        if n in by_id:
            problems.append(f"{where}: Unique TMA sample ID {n} repeated")
            continue
        if not 1 <= n <= n_layout:
            problems.append(f"{where}: Unique TMA sample ID {n} outside 1..{n_layout}")
            continue
        by_id[n] = rec
    missing = sorted(set(range(1, n_layout + 1)) - set(by_id))
    if missing:
        problems.append(f"Lookup: layout positions {missing} have no row")
    grid_occ = {position_id(x, y, rows): (x, y) for (x, y) in ids}
    donor_occ = {position_id(x, y, rows): v for (x, y), v in donors.items()}
    look_occ = {n: r["donor block id"] for n, r in by_id.items() if r.get("donor block id")}
    for n in sorted(set(grid_occ) | set(donor_occ) | set(look_occ)):
        where = ", ".join(s for s, ok in ((block, n in grid_occ), (f"{block}_Donor", n in donor_occ),
                                           ("Lookup", n in look_occ)) if ok)
        if not (n in grid_occ and n in donor_occ and n in look_occ):
            problems.append(f"position {n}: occupied only in {where}")
        elif donor_occ[n] != look_occ[n]:
            problems.append(f"position {n}: donor {donor_occ[n]!r} in {block}_Donor, {look_occ[n]!r} in Lookup")
    for key in ("number of core", "number of donor"):
        if info.get(key, "") and _int(info.get(key)) is None:
            problems.append(f"Information: {key.title()} {info.get(key)!r} is not a whole number")
    declared_cores = _int(info.get("number of core"))
    if declared_cores is not None and declared_cores != len(look_occ):
        problems.append(f"Information: Number of Core {declared_cores}, {len(look_occ)} occupied positions")
    declared_donors = _int(info.get("number of donor"))
    n_donors = len(set(look_occ.values()))
    if declared_donors is not None and declared_donors != n_donors:
        problems.append(f"Information: Number of Donor {declared_donors}, {n_donors} distinct donor blocks")

    # ---- optional field: Spot Slide ----
    has_spot = "spot slide" in header
    positions: Dict[int, Position] = {}
    for n in range(1, n_layout + 1):
        x, y = (n - 1) // rows + 1, (n - 1) % rows + 1
        rec = by_id.get(n, {})
        donor = rec.get("donor block id", "")
        spot = rec.get("spot slide", "") if has_spot else ""
        if spot and not donor:
            problems.append(f"Lookup!R{rec['_row']}: Spot Slide {spot!r} on the empty position {n}")
        elif spot and donor:
            m = SPOT.match(spot)
            if not m:
                problems.append(f"Lookup!R{rec['_row']}: Spot Slide {spot!r} unreadable")
            else:
                sc, sr, sx, sd = int(m.group(1)), int(m.group(2)), int(m.group(3)), m.group(4)
                if (sc, sr, sx) != (x, y, x) or sd != donor:
                    problems.append(f"Lookup!R{rec['_row']}: Spot Slide {spot!r} does not match "
                                    f"position {n} (X{x} Y{y}, donor {donor!r})")
        positions[n] = Position(n, x, y, donor, rec.get("unique tma sample id", "") if rec else "",
                                rec.get("donor core height", ""), rec.get("annotation name", ""), spot)
    if problems:
        raise MapSheetError(path, problems)
    diameter = info.get("diameter", "")
    try:
        diameter_mm = float(diameter) if diameter else None
    except ValueError:
        diameter_mm = None
    return MapSheet(path=os.path.abspath(path), sha256=file_sha256(path), size=os.path.getsize(path),
                    block_id=block, block_name=info.get("tma block name", ""), rows=rows, cols=cols,
                    layout_name=info.get("layout name", ""), diameter_mm=diameter_mm,
                    declared_cores=declared_cores, declared_donors=declared_donors,
                    positions=positions, has_spot_slide=has_spot)
